"""--juntar (junção interna por chave) e --so-se-atualizou (execução agendada)."""
from __future__ import annotations

import openpyxl
import pytest
import responses

from conftest import tg
from test_faixas import por_igualdade

BENEF = [{"id_beneficiario": i, "uf_beneficiario": uf, "nome_beneficiario": f"MUNICIPIO {i}",
          "cnpj_beneficiario": str(i), "id_ente": i} for i, uf in [(1, "SP"), (2, "SP"), (3, "RJ"), (4, "MG")]]
PLANOS = [{"id_plano_acao": 100 + i, "ano_plano_acao": 2025, "id_beneficiario": b, "valor_custeio_plano_acao": 10.0 * i}
          for i, b in enumerate([1, 1, 2, 3, 3, 3, 4, 9])]
ARGS = ["extrair", "--modulo", "especiais", "--tabela", "planos-acao-especiais", "--tabela", "beneficiarios-especiais",
        "-f", "beneficiarios-especiais:uf_beneficiario=SP", "--juntar", "id_beneficiario"]


def _servir(servir_tabela):
    ch_b = servir_tabela("especiais", "beneficiarios-especiais", BENEF, filtro=por_igualdade)
    ch_p = servir_tabela("especiais", "planos-acao-especiais", PLANOS, filtro=por_igualdade)
    return ch_b, ch_p


def test_juncao_interna_gera_aba_juntado_e_restringe_as_tabelas(servir_tabela, tmp_path):
    _servir(servir_tabela)
    assert tg.main(ARGS + ["--saida", str(tmp_path), "--arquivo", "pix_sp"]) == 0
    wb = openpyxl.load_workbook(tmp_path / "pix_sp.xlsx")
    assert wb.sheetnames == ["juntado", "planos-acao-especiais", "beneficiarios-especiais", "metadados"]
    cab = [c.value for c in wb["juntado"][1]]
    assert "uf_beneficiario" in cab and "valor_custeio_plano_acao" in cab and cab.count("id_beneficiario") == 1
    ufs = {r[cab.index("uf_beneficiario")] for r in wb["juntado"].iter_rows(min_row=2, values_only=True)}
    assert ufs == {"SP"} and wb["juntado"].max_row - 1 == 3          # planos 100, 101 (benef 1) e 102 (benef 2)
    assert wb["planos-acao-especiais"].max_row - 1 == 3             # restrita às chaves que casaram
    meta = [r for r in wb["metadados"].iter_rows(values_only=True)]
    assert any(r[0] == "juntado" and r[1] == "juncao" and "id_beneficiario" in r[2] for r in meta)


def test_juncao_sem_filtros_mantem_abas_completas(servir_tabela, tmp_path):
    """'Uma aba por tabela' sem filtro: cada aba vem inteira; só a aba juntado é a interseção."""
    _servir(servir_tabela)
    args = ["extrair", "--modulo", "especiais", "--tabela", "planos-acao-especiais",
            "--tabela", "beneficiarios-especiais", "--juntar", "id_beneficiario",
            "--saida", str(tmp_path), "--arquivo", "x"]
    assert tg.main(args) == 0
    wb = openpyxl.load_workbook(tmp_path / "x.xlsx")
    assert wb["planos-acao-especiais"].max_row - 1 == len(PLANOS)       # inclui o plano do benef. 9 (sem par)
    assert wb["beneficiarios-especiais"].max_row - 1 == len(BENEF)
    assert wb["juntado"].max_row - 1 == 7                               # 8 planos, 1 sem beneficiário


def test_juncao_com_filtro_numa_tabela_restringe_so_as_outras(servir_tabela, tmp_path):
    _servir(servir_tabela)
    assert tg.main(ARGS + ["--saida", str(tmp_path), "--arquivo", "x"]) == 0
    wb = openpyxl.load_workbook(tmp_path / "x.xlsx")
    assert wb["beneficiarios-especiais"].max_row - 1 == 2   # os 2 de SP, pelo filtro do próprio usuário
    assert wb["planos-acao-especiais"].max_row - 1 == 3     # restrita aos de SP pela junção


def test_juncao_com_poucas_chaves_consulta_a_outra_tabela_por_chave(servir_tabela, tmp_path):
    ch_b, ch_p = _servir(servir_tabela)
    # página de 1 linha: baixar planos inteiro = 8 páginas > 2 chaves → consulta por chave
    assert tg.main(ARGS + ["--saida", str(tmp_path), "--arquivo", "x", "--tamanho-pagina", "1"]) == 0
    assert ch_b[0]["uf_beneficiario"] == "SP"                      # a tabela filtrada vai primeiro
    assert {c["id_beneficiario"] for c in ch_p if "id_beneficiario" in c} == {"1", "2"}
    assert len([c for c in ch_p if "id_beneficiario" not in c]) == 1   # só a contagem para estimar
    wb = openpyxl.load_workbook(tmp_path / "x.xlsx")
    meta = {(r[0], r[1]): r[2] for r in wb["metadados"].iter_rows(values_only=True)}
    assert "2 valores" in meta[("planos-acao-especiais", "filtros")]  # sem despejar a lista de ids


def test_juncao_com_muitas_chaves_baixa_a_tabela_inteira(servir_tabela, tmp_path):
    _, ch_p = _servir(servir_tabela)
    assert tg.main(ARGS + ["--saida", str(tmp_path), "--arquivo", "x"]) == 0   # 1 página de 200 < 2 chaves
    assert all("id_beneficiario" not in c for c in ch_p if c["tamanho_da_pagina"] != "1")


def test_juncao_sem_chaves_em_comum_nao_baixa_a_segunda_tabela(servir_tabela, tmp_path):
    _, ch_p = _servir(servir_tabela)
    args = [a if a != "beneficiarios-especiais:uf_beneficiario=SP" else "beneficiarios-especiais:uf_beneficiario=AC"
            for a in ARGS]
    assert tg.main(args + ["--saida", str(tmp_path), "--arquivo", "x"]) == 0
    assert not [c for c in ch_p if c["tamanho_da_pagina"] != "1"]
    assert openpyxl.load_workbook(tmp_path / "x.xlsx")["juntado"].max_row == 1


def test_juncao_chave_inexistente_e_tabela_unica(api, tmp_path, capsys):
    base = ARGS[:-1]
    assert tg.main(base + ["id_ente", "--saida", str(tmp_path)]) == 1
    assert "id_ente" in capsys.readouterr().err
    assert tg.main(["extrair", "--modulo", "especiais", "--tabela", "beneficiarios-especiais",
                    "--juntar", "id_beneficiario", "--saida", str(tmp_path)]) == 1


# ---------- --so-se-atualizou ----------

def _args_diarios(tmp_path, formato="xlsx"):
    return ["extrair", "--modulo", "especiais", "--tabela", "beneficiarios-especiais", "-f", "uf_beneficiario=SP",
            "--saida", str(tmp_path), "--arquivo", "diario", "--formato", formato, "--so-se-atualizou"]


@pytest.mark.parametrize("formato", ["xlsx", "csv"])
def test_so_se_atualizou_pula_quando_base_nao_mudou(servir_tabela, tmp_path, capsys, formato):
    ch = servir_tabela("especiais", "beneficiarios-especiais", BENEF, filtro=por_igualdade)
    assert tg.main(_args_diarios(tmp_path, formato)) == 0
    n = len(ch)
    capsys.readouterr()
    assert tg.main(_args_diarios(tmp_path, formato)) == 0
    out, err = capsys.readouterr()
    assert len(ch) == n                                  # não baixou de novo
    assert "diario" in out and "sem atualização" in err


def test_so_se_atualizou_extrai_quando_base_mudou(api, servir_tabela, tmp_path):
    ch = servir_tabela("especiais", "beneficiarios-especiais", BENEF, filtro=por_igualdade)
    assert tg.main(_args_diarios(tmp_path)) == 0
    n = len(ch)
    url = f"{tg.BASE}/especiais/data-atualizacao"
    api.remove(responses.GET, url)
    api.get(url, json={"data_ultima_atualizacao": "2026-10-16T00:00:00"})
    assert tg.main(_args_diarios(tmp_path)) == 0
    assert len(ch) > n


def test_so_se_atualizou_extrai_quando_filtros_mudaram(servir_tabela, tmp_path):
    ch = servir_tabela("especiais", "beneficiarios-especiais", BENEF, filtro=por_igualdade)
    assert tg.main(_args_diarios(tmp_path)) == 0
    n = len(ch)
    args = [a if a != "uf_beneficiario=SP" else "uf_beneficiario=RJ" for a in _args_diarios(tmp_path)]
    assert tg.main(args) == 0
    assert len(ch) > n


def test_so_se_atualizou_exige_arquivo_fixo(api, tmp_path, capsys):
    args = [a for a in _args_diarios(tmp_path) if a not in ("--arquivo", "diario")]
    assert tg.main(args) == 1
    assert "--arquivo" in capsys.readouterr().err
