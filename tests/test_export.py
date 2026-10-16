"""Exportação xlsx/csv e CLI ponta a ponta."""
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from conftest import ler_fixture, tg
from test_paginacao import linhas_benef


def _servir_programas(servir_tabela):
    linhas = [l for p in (1, 2, 3) for l in ler_fixture(f"pagina_especiais_programas-especiais_p{p}_t6.json")["data"]]
    servir_tabela("especiais", "programas-especiais", linhas)
    return linhas


def test_xlsx_com_aba_por_tabela_e_metadados(servir_tabela, tmp_path):
    _servir_programas(servir_tabela)
    servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(4))
    r1 = tg.extrair("especiais", "programas-especiais", ["ano_programa=2026"])
    r2 = tg.extrair("especiais", "beneficiarios-especiais", [])
    [p] = tg.exportar([r1, r2], "xlsx", tmp_path, "saida_teste")
    assert p == tmp_path / "saida_teste.xlsx"
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["programas-especiais", "beneficiarios-especiais", "metadados"]
    ws = wb["programas-especiais"]
    assert ws.max_row == 16  # cabeçalho + 15
    cab = [c.value for c in ws[1]]
    i = cab.index("data_inicio_ciencia_programa") + 1
    assert isinstance(ws.cell(2, i).value, dt.datetime)
    assert ws.cell(2, i).number_format == "DD/MM/YYYY"
    meta = [[c.value for c in row] for row in wb["metadados"].iter_rows()]
    assert meta[0] == ["tabela", "chave", "valor"]
    assert ["programas-especiais", "filtros", "ano_programa=2026"] in meta
    assert ["beneficiarios-especiais", "linhas", 4] in meta
    assert not list(tmp_path.glob("*.tmp*"))


def test_xlsx_divide_em_abas_acima_do_limite(servir_tabela, tmp_path, monkeypatch):
    monkeypatch.setattr(tg, "EXCEL_MAX_ROWS", 3)
    servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(7))
    r = tg.extrair("especiais", "beneficiarios-especiais", [])
    [p] = tg.exportar([r], "xlsx", tmp_path, "grande")
    wb = openpyxl.load_workbook(p)
    abas = [n for n in wb.sheetnames if n != "metadados"]
    assert len(abas) == 3 and all(len(n) <= 31 for n in abas)
    assert [wb[n].max_row - 1 for n in abas] == [3, 3, 1]


def test_nome_de_aba_longo_e_truncado_em_31(api, tmp_path, servir_tabela):
    servir_tabela("especiais", "ordens-pagamentos-ordens-bancarias-especiais", [])
    r = tg.extrair("especiais", "ordens-pagamentos-ordens-bancarias-especiais", [])
    [p] = tg.exportar([r], "xlsx", tmp_path, "x")
    assert all(len(n) <= 31 for n in openpyxl.load_workbook(p).sheetnames)


def test_csv_ptbr(servir_tabela, tmp_path):
    _servir_programas(servir_tabela)
    r = tg.extrair("especiais", "programas-especiais", [])
    p, pm = tg.exportar([r], "csv", tmp_path, "prog")
    assert pm.name == "prog_metadados.csv"
    assert "programas-especiais;filtros;(nenhum)" in pm.read_text(encoding="utf-8-sig")
    bruto = p.read_bytes()
    assert bruto.startswith(b"\xef\xbb\xbf")  # BOM: Excel abre em UTF-8
    txt = bruto.decode("utf-8-sig")
    cab, l1 = txt.splitlines()[:2]
    assert ";" in cab and "," not in cab
    assert "6779708728,64" in l1          # vírgula decimal
    assert "30/03/2026" in l1             # data pt-BR
    assert "Ministério" in txt


def test_csv_uma_saida_por_tabela(servir_tabela, tmp_path):
    _servir_programas(servir_tabela)
    servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(2))
    rs = [tg.extrair("especiais", "programas-especiais", []), tg.extrair("especiais", "beneficiarios-especiais", [])]
    ps = tg.exportar(rs, "csv", tmp_path, "lote")
    assert [p.name for p in ps] == ["lote_programas-especiais.csv", "lote_beneficiarios-especiais.csv",
                                    "lote_metadados.csv"]


# ---------- CLI ----------

def test_cli_extrair_imprime_so_o_caminho_no_stdout(servir_tabela, tmp_path, capsys):
    servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(3))
    rc = tg.main(["extrair", "--modulo", "especiais", "--tabela", "beneficiarios-especiais",
                  "--saida", str(tmp_path), "--arquivo", "fixo"])
    out, err = capsys.readouterr()
    assert rc == 0
    assert out.strip() == str(tmp_path / "fixo.xlsx")
    assert "beneficiarios-especiais" in err


def test_cli_arquivo_fixo_sobrescreve(servir_tabela, tmp_path):
    servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(3))
    args = ["extrair", "--modulo", "especiais", "--tabela", "beneficiarios-especiais",
            "--saida", str(tmp_path), "--arquivo", "fixo"]
    assert tg.main(args) == 0 and tg.main(args) == 0
    assert [p.name for p in tmp_path.iterdir()] == ["fixo.xlsx"]


def test_cli_erro_vai_para_stderr_com_codigo_1(api, capsys):
    rc = tg.main(["contar", "--modulo", "especiais", "--tabela", "nao-existe"])
    out, err = capsys.readouterr()
    assert rc == 1 and out == "" and "beneficiarios-especiais" in err


def test_cli_estrito_retorna_3_quando_ha_aviso(servir_tabela, tmp_path):
    servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(4), total=5)
    base = ["extrair", "--modulo", "especiais", "--tabela", "beneficiarios-especiais", "--saida", str(tmp_path)]
    assert tg.main(base) == 0
    assert tg.main(base + ["--estrito"]) == 3


def test_cli_varias_tabelas_filtro_com_escopo(servir_tabela, tmp_path):
    """Filtro sem escopo precisa valer para todas as tabelas; com 'tabela:' vale só para ela."""
    ch_b = servir_tabela("especiais", "beneficiarios-especiais", linhas_benef(3))
    _servir_programas(servir_tabela)
    base = ["extrair", "--modulo", "especiais", "--tabela", "beneficiarios-especiais",
            "--tabela", "programas-especiais", "--saida", str(tmp_path), "--arquivo", "j"]
    assert tg.main(base + ["-f", "uf_beneficiario=SP"]) == 1  # programas não tem esse filtro
    assert tg.main(base + ["-f", "beneficiarios-especiais:uf_beneficiario=SP"]) == 0
    assert ch_b[-1]["uf_beneficiario"] == "SP"
    wb = openpyxl.load_workbook(tmp_path / "j.xlsx")
    assert wb.sheetnames == ["beneficiarios-especiais", "programas-especiais", "metadados"]


def test_cli_colunas_mostra_colunas_e_filtros(api, capsys):
    assert tg.main(["colunas", "--modulo", "ted", "--tabela", "planos-acao-pareceres"]) == 0
    out = capsys.readouterr().out
    assert "id_parecer" in out and "PK" in out.upper() and "dt_data_parecer" in out


@pytest.mark.parametrize("cmd", ["tabelas"])
def test_cli_tabelas(api, capsys, cmd):
    assert tg.main([cmd, "--modulo", "ted"]) == 0
    assert "planos-acao-pareceres" in capsys.readouterr().out.split()
