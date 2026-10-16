"""Módulo `discricionarias`: convênios e contratos de repasse (antigo Siconv), via CSVs do /downloads.

Não há OpenAPI: a lista de arquivos vem da listagem do container; as colunas, do cabeçalho do CSV; os tipos,
da convenção de nomes do modelo oficial (DIA_/DATA_/DT_ = data, VL_/VALOR_ = número, QTD_/ANO/MES = inteiro).
"""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pandas as pd
import pytest

from conftest import tg

LISTA = f"{tg.BASE}/downloads/dadosgov/"
MOD = "Thu, 15 Oct 2026 11:09:49 GMT"

CONVENIO = (
    "NR_CONVENIO;ID_PROPOSTA;DIA;ANO;DIA_ASSIN_CONV;SIT_CONVENIO;VL_GLOBAL_CONV;QTD_TA\n"
    "900001;10;5;2025;05/03/2025;Em execução;1234,56;1\n"
    "900002;11;7;2025;07/04/2025;Prestação de Contas Concluída;99,9;\n"
    "900003;12;9;2024;09/12/2024;Em execução;500000;2\n"
    "000004;13;;;;Proposta;0;\n"
)
PROPOSTA = (
    "ID_PROPOSTA;UF_PROPONENTE;DESC_ORGAO_SUP;DIA_PROP;MES_PROP;ANO_PROP;DIA_PROPOSTA\n"
    "10;SP;MINISTÉRIO DA SAÚDE;5;6;2009;05/06/2009\n"
    "11;RJ;MINISTERIO DA EDUCACAO;5;3;2010;05/03/2010\n"
    "12;SP;MINISTÉRIO DA SAÚDE;5;7;2011;05/07/2011\n"
)


def _zip(nome: str, csv: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(f"{nome}.csv", ("﻿" + csv).encode("utf-8"))
    return buf.getvalue()


def _listagem(blobs: dict[str, bytes]) -> str:
    itens = "".join(
        f"<Blob><Name>{n}</Name><Properties><Last-Modified>{MOD}</Last-Modified>"
        f"<Etag>0xETAG{len(b)}</Etag><Content-Length>{len(b)}</Content-Length></Properties></Blob>"
        for n, b in blobs.items())
    return f'<?xml version="1.0" encoding="utf-8"?><EnumerationResults><Blobs>{itens}</Blobs><NextMarker/></EnumerationResults>'


@pytest.fixture
def downloads(api, tmp_path, monkeypatch):
    monkeypatch.setattr(tg, "CACHE_DIR", tmp_path / "cache")
    tg._cache_listagem.clear()
    tg._cache_csv.clear()
    blobs = {"siconv_convenio.zip": _zip("siconv_convenio", CONVENIO),
             "siconv_proposta.zip": _zip("siconv_proposta", PROPOSTA),
             "data_carga_siconv.zip": _zip("data_carga_siconv", "DATA\n15/10/2026\n")}
    api.get(LISTA, body=_listagem(blobs), content_type="application/xml")
    for n, b in blobs.items():
        api.get(LISTA + n, body=b, content_type="application/octet-stream")
    return api


def _baixados(api, nome):
    return sum(c.request.url.endswith(nome) for c in api.calls)


def test_tabelas_vem_da_listagem(downloads):
    assert tg.tabelas("discricionarias") == ["convenio", "proposta"]
    assert tg.tabelas("convenios") == ["convenio", "proposta"]  # apelido


def test_tipos_pela_convencao_de_nomes(downloads):
    ep = tg.endpoint("discricionarias", "convenio")
    tipos = {c.nome: (c.tipo, c.formato) for c in ep.colunas}
    assert tipos["DIA_ASSIN_CONV"] == ("string", "date")
    assert tipos["VL_GLOBAL_CONV"] == ("number", None)
    assert tipos["QTD_TA"] == ("integer", None)
    assert tipos["ANO"] == ("integer", None) and tipos["DIA"] == ("integer", None)
    assert tipos["NR_CONVENIO"] == ("string", None)       # zeros à esquerda preservados
    tp = {c.nome: (c.tipo, c.formato) for c in tg.endpoint("discricionarias", "proposta").colunas}
    assert tp["DIA_PROP"] == ("integer", None)            # trio DIA_/MES_/ANO_PROP: número do dia
    assert tp["DIA_PROPOSTA"] == ("string", "date")


def test_extrair_tipa_e_filtra(downloads):
    r = tg.extrair("discricionarias", "convenio", ["ANO=2025"])
    assert list(r.df["NR_CONVENIO"]) == ["900001", "900002"]
    assert r.df["VL_GLOBAL_CONV"].iloc[0] == pytest.approx(1234.56)
    assert r.df["DIA_ASSIN_CONV"].iloc[0] == pd.Timestamp("2025-03-05")
    assert pd.isna(r.df["QTD_TA"].iloc[1])
    assert r.meta["data_atualizacao_base"] == MOD
    assert r.avisos == []


def test_texto_compara_contem_sem_acento_e_maiusculas(downloads):
    r = tg.extrair("discricionarias", "convenio", ["SIT_CONVENIO=EXECUCAO"])
    assert list(r.df["NR_CONVENIO"]) == ["900001", "900003"]


def test_codigos_comparam_por_igualdade_exata(downloads):
    """ID_/COD_/NR_/CNPJ... são códigos: ID_PROPOSTA=1 não pode trazer 10, 11, 12."""
    assert tg.contar("discricionarias", "convenio", ["ID_PROPOSTA=1"]) == 0
    r = tg.extrair("discricionarias", "convenio", ["NR_CONVENIO=000004"])
    assert list(r.df["NR_CONVENIO"]) == ["000004"]


def test_faixas_de_valor_e_data(downloads):
    r = tg.extrair("discricionarias", "convenio", ["VL_GLOBAL_CONV>=1000", "DIA_ASSIN_CONV<=2025-03-31"])
    assert list(r.df["NR_CONVENIO"]) == ["900001", "900003"]


def test_contar_bate_com_extrair_e_baixa_uma_vez(downloads):
    assert tg.contar("discricionarias", "convenio", ["ANO=2025"]) == 2
    tg.extrair("discricionarias", "convenio", ["ANO=2025"])
    assert _baixados(downloads, "siconv_convenio.zip") == 1   # cache local do zip


def test_filtro_inexistente_lista_colunas(downloads):
    with pytest.raises(tg.ErroTG, match="DIA_ASSIN_CONV"):
        tg.contar("discricionarias", "convenio", ["ano_convenio=2025"])


def test_arquivo_inexistente_lista_validos(downloads):
    with pytest.raises(tg.ErroTG, match="convenio"):
        tg.endpoint("discricionarias", "convenios_saude")


def test_juncao_convenio_proposta(downloads, tmp_path):
    rc = tg.main(["extrair", "--modulo", "discricionarias", "--tabela", "convenio", "--tabela", "proposta",
                  "-f", "proposta:DESC_ORGAO_SUP=saude", "-f", "convenio:ANO=2025",
                  "--juntar", "ID_PROPOSTA", "--saida", str(tmp_path), "--arquivo", "conv"])
    assert rc == 0
    wb = openpyxl.load_workbook(tmp_path / "conv.xlsx")
    assert wb.sheetnames == ["juntado", "convenio", "proposta", "metadados"]
    cab = [c.value for c in wb["juntado"][1]]
    linhas = list(wb["juntado"].iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 1 and linhas[0][cab.index("NR_CONVENIO")] == "900001"


def test_valor_fora_do_formato_gera_aviso(api, downloads):
    ruim = CONVENIO + "900005;14;1;2025;31/02/2025;Em execução;abc;1\n"
    blobs = {"siconv_convenio.zip": _zip("siconv_convenio", ruim)}
    tg._cache_listagem.clear()
    api.replace("GET", LISTA, body=_listagem(blobs), content_type="application/xml")
    api.replace("GET", LISTA + "siconv_convenio.zip", body=blobs["siconv_convenio.zip"])
    r = tg.extrair("discricionarias", "convenio", ["NR_CONVENIO=900005"])
    assert any("DIA_ASSIN_CONV" in a for a in r.avisos) and any("VL_GLOBAL_CONV" in a for a in r.avisos)
